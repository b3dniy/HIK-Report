# -*- coding: utf-8 -*-
"""
HIK-Report — отчёт по входам пользователей для регистраторов Hikvision (ISAPI).
Всё в терминале, без открытия веб-интерфейса; итог — оформленный Excel-файл.

Регистраторы задаются пользователем в меню настроек («0») и сохраняются
внутри самого exe (NTFS ADS). В коде никаких адресов по умолчанию нет.
"""

import sys
import os
import time
import json
import math
import colorsys
import socket
import getpass
import threading
import subprocess
import datetime as dt
import xml.etree.ElementTree as ET
import uuid


# ------------------------------------------------------------------ #
#  Автоустановка зависимостей
# ------------------------------------------------------------------ #
def _ensure(pkg, import_name=None):
    import_name = import_name or pkg
    try:
        __import__(import_name)
    except ImportError:
        print(f"  Устанавливаю зависимость: {pkg} ...")
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", "--quiet", pkg]
        )


# В собранном .exe зависимости уже внутри — ставить ничего не нужно.
if not getattr(sys, "frozen", False):
    for _p, _i in (("requests", "requests"),
                   ("openpyxl", "openpyxl"),
                   ("colorama", "colorama")):
        _ensure(_p, _i)

import requests
from requests.auth import HTTPDigestAuth, HTTPBasicAuth
from colorama import init as colorama_init, Fore, Style
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def _enable_vt():
    """Включает ANSI/VT в консоли Windows (для 24-битного цвета). True если ок."""
    if os.name != "nt":
        return True
    try:
        import ctypes
        k = ctypes.windll.kernel32
        ok = False
        for handle_id in (-11, -12):          # STDOUT, STDERR
            h = k.GetStdHandle(handle_id)
            mode = ctypes.c_uint32()
            if k.GetConsoleMode(h, ctypes.byref(mode)):
                k.SetConsoleMode(h, mode.value | 0x0004)  # VT processing
                ok = True
        return ok
    except Exception:
        return False


TRUECOLOR = _enable_vt()
if TRUECOLOR:
    colorama_init(strip=False, convert=False)  # пропускаем ANSI как есть (24-бит)
else:
    colorama_init()
requests.packages.urllib3.disable_warnings()  # на случай https с самоподписанным


# ------------------------------------------------------------------ #
#  Консоль в UTF-8 (для русского текста и псевдографики)
# ------------------------------------------------------------------ #
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stdin.reconfigure(encoding="utf-8")
except Exception:
    pass


def output_dir():
    """Куда сохранять отчёт: рядом с программой, иначе на Рабочий стол."""
    if getattr(sys, "frozen", False):
        base = os.path.dirname(sys.executable)
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    try:
        t = os.path.join(base, ".write_test")
        open(t, "w").close()
        os.remove(t)
        return base
    except Exception:
        desktop = os.path.join(os.path.expanduser("~"), "Desktop")
        return desktop if os.path.isdir(desktop) else os.path.expanduser("~")


# ------------------------------------------------------------------ #
#  Конфигурация регистраторов
# ------------------------------------------------------------------ #
HTTP_TIMEOUT = 8

# По умолчанию список пуст: пользователь добавляет свои регистраторы
# через меню настроек («0»). Настройки сохраняются внутри самого exe (ADS).
DEFAULT_RECORDERS = []


def _cfg_stream():
    """Путь к хранилищу настроек — альтернативный поток NTFS внутри самого exe.
    Отдельного видимого файла нет; настройки «прилипают» к exe."""
    base = sys.executable if getattr(sys, "frozen", False) \
        else os.path.abspath(__file__)
    return base + ":hikcfg"


def load_recorders():
    """Читает регистраторы из ADS exe; если пусто/ошибка — значения по умолчанию."""
    try:
        with open(_cfg_stream(), "r", encoding="utf-8") as f:
            data = json.load(f)
        out = []
        for r in (data.get("recorders") or []):
            name = str(r.get("name", "")).strip() or str(r.get("lan_ip", ""))
            out.append({
                "name": name,
                "lan_ip": str(r.get("lan_ip", "")).strip(),
                "lan_port": int(r.get("lan_port", 80) or 80),
                "rem_host": str(r.get("rem_host", "")).strip(),
                "rem_port": int(r.get("rem_port", 0) or 0),
            })
        if out:
            return out
    except Exception:
        pass
    return [dict(r) for r in DEFAULT_RECORDERS]


def save_recorders(recs):
    """Сохраняет регистраторы в ADS exe. Возвращает (ok, сообщение об ошибке)."""
    try:
        with open(_cfg_stream(), "w", encoding="utf-8") as f:
            json.dump({"recorders": recs}, f, ensure_ascii=False, indent=2)
        return True, ""
    except Exception as e:
        return False, str(e)


# ------------------------------------------------------------------ #
#  Вспомогательные функции вывода
# ------------------------------------------------------------------ #
def cls():
    try:
        CON.clear_badges()
    except Exception:
        pass
    os.system("cls" if os.name == "nt" else "clear")


def banner(text):
    line = "═" * (len(text) + 2)
    print(Fore.CYAN + "╔" + line + "╗")
    print(Fore.CYAN + "║ " + Style.BRIGHT + text + Style.NORMAL + Fore.CYAN + " ║")
    print(Fore.CYAN + "╚" + line + "╝" + Style.RESET_ALL)


def progress(cur, total, prefix="", width=34):
    total = max(total, 1)
    frac = min(cur / total, 1.0)
    filled = int(width * frac)
    bar = "█" * filled + "░" * (width - filled)
    sys.stdout.write(
        "\r" + Fore.GREEN + prefix.ljust(26) + Style.RESET_ALL +
        " |" + Fore.CYAN + bar + Style.RESET_ALL + f"| {int(frac*100):3d}%"
    )
    sys.stdout.flush()
    if cur >= total:
        sys.stdout.write("\n")
        sys.stdout.flush()


_START_HUE = 0.78   # фиолетовый — стартовый цвет радуги


def _rainbow_rgb(h):
    """Полная радуга из HSV. h — любое число, берётся дробная часть."""
    r, g, b = colorsys.hsv_to_rgb(h % 1.0, 1.0, 1.0)
    return int(r * 255), int(g * 255), int(b * 255)


_BAR_W = 34


class _Console:
    """Прокси stdout + один фоновый поток анимации.

    • снизу рисуется радужный прогресс-бар;
    • выше — бейджи SUCCESS/ERROR, которые ПРОДОЛЖАЮТ мигать на следующих
      шагах, пока их строка видна на экране (ушла в прокрутку — застывает).
    Процесс не задерживается: всё рисует фоновый поток. Позиция курсора
    сохраняется/восстанавливается (ESC 7 / ESC 8), поэтому даже ввод с
    клавиатуры не портится."""

    def __init__(self, real):
        self.real = real
        self.lock = threading.RLock()
        self.badges = []        # [{off, on, off_str}]
        self.bar = None         # {prefix, frac, phase} | None
        self.frame = 0
        self._stop = threading.Event()
        self._thr = None

    # ---- интерфейс потока вывода (sys.stdout) ----
    def write(self, s):
        if not s:
            return 0
        with self.lock:
            n = s.count("\n")
            if n and self.badges:
                for b in self.badges:
                    b["off"] += n
            self.real.write(s)
            self.real.flush()
            if n:
                self._drop_scrolled()
        return len(s)

    def flush(self):
        try:
            self.real.flush()
        except Exception:
            pass

    def isatty(self):
        try:
            return self.real.isatty()
        except Exception:
            return False

    def fileno(self):
        return self.real.fileno()

    def __getattr__(self, name):
        return getattr(self.real, name)

    # ---- служебное ----
    def _term_lines(self):
        try:
            return os.get_terminal_size().lines
        except Exception:
            return 25

    def _drop_scrolled(self):
        if self.badges:
            h = self._term_lines()
            self.badges = [b for b in self.badges if b["off"] < h - 1]

    def start(self):
        if TRUECOLOR and self._thr is None:
            self._thr = threading.Thread(target=self._run, daemon=True)
            self._thr.start()
        return self

    def stop(self):
        self._stop.set()
        if self._thr:
            self._thr.join(timeout=0.5)
        try:
            self.real.write("\x1b[0m")
            self.real.flush()
        except Exception:
            pass

    def _run(self):
        while not self._stop.is_set():
            self._tick()
            time.sleep(0.035)

    def _bar_str(self):
        bar = self.bar
        if bar is None:
            return None
        bar["phase"] = (bar["phase"] + 0.015) % 1.0
        frac, prefix, phase = bar["frac"], bar["prefix"], bar["phase"]
        filled = int(round(_BAR_W * frac))
        cells = []
        for i in range(filled):
            hL = _START_HUE - i / _BAR_W + phase
            hR = _START_HUE - (i + 0.5) / _BAR_W + phase
            rL, gL, bL = _rainbow_rgb(hL)
            rR, gR, bR = _rainbow_rgb(hR)
            cells.append(f"\x1b[38;2;{rL};{gL};{bL};48;2;{rR};{gR};{bR}m▌")
        tail = "\x1b[0m"
        if filled < _BAR_W:
            tail += "\x1b[38;2;55;58;64m" + "░" * (_BAR_W - filled) + "\x1b[0m"
        pref = f"\x1b[92m{prefix:<26}\x1b[0m"
        return f"\r{pref} |{''.join(cells)}{tail}| {int(frac * 100):3d}%"

    def _tick(self):
        with self.lock:
            self.frame += 1
            blink_on = (self.frame // 12) % 2 == 0     # смена ~каждые 0.42 c
            h = self._term_lines()
            out = []
            for b in self.badges:
                d = b["off"]
                if d <= 0 or d >= h - 1:
                    continue
                line = b["on"] if blink_on else b["off_str"]
                out.append(f"\x1b7\x1b[{d}A\r{line}\x1b[K\x1b8")
            bar = self._bar_str()
            if bar is not None:
                out.append(bar)
            if out:
                try:
                    self.real.write("".join(out))
                    self.real.flush()
                except Exception:
                    pass

    # ---- API ----
    def bar_start(self, prefix):
        with self.lock:
            self.bar = {"prefix": prefix, "frac": 0.0, "phase": 0.0}

    def bar_set(self, frac, prefix=None):
        with self.lock:
            if self.bar is None:
                return
            self.bar["frac"] = 0.0 if frac < 0 else (1.0 if frac > 1 else frac)
            if prefix is not None:
                self.bar["prefix"] = prefix

    def _push_badge(self, on, off_str, plain):
        if TRUECOLOR:
            self.real.write("\r\x1b[K" + on + "\n")
            self.real.flush()
            for bd in self.badges:
                bd["off"] += 1
            self.badges.append({"off": 1, "on": on, "off_str": off_str})
            self._drop_scrolled()
        else:
            self.real.write("\r" + plain + "\n")
            self.real.flush()

    def bar_finish(self, ok=True, prefix=None):
        with self.lock:
            b = self.bar
            self.bar = None
            prefix = prefix or (b["prefix"] if b else "")
            word = "✓  SUCCESS" if ok else "✗  ERROR"
            field = word.center(_BAR_W)
            cmain = "\x1b[1;92m" if ok else "\x1b[1;91m"
            pref = f"\x1b[92m{prefix:<26}\x1b[0m"
            on = f"{pref} |{cmain}{field}\x1b[0m|"
            off = f"{pref} |\x1b[1;97m{field}\x1b[0m|"
            self._push_badge(on, off, f"{prefix:<26} |{field}|")

    def add_badge(self, text, ok=True):
        with self.lock:
            cmain = "\x1b[1;92m" if ok else "\x1b[1;91m"
            on = f"{cmain}{text}\x1b[0m"
            off = f"\x1b[1;97m{text}\x1b[0m"
            self._push_badge(on, off, text)

    def clear_badges(self):
        """Сброс отслеживаемых бейджей/бара (перед очисткой экрана)."""
        with self.lock:
            self.badges = []
            self.bar = None


class ProgressBar:
    """Тонкая обёртка: сам бар рисует фоновый поток _Console (без задержек)."""

    def __init__(self, prefix="", width=_BAR_W):
        self.prefix = prefix

    def start(self):
        CON.bar_start(self.prefix)
        if not TRUECOLOR:
            sys.stdout.write(f"{self.prefix:<26} |{'░' * _BAR_W}|   0%")
            sys.stdout.flush()
        return self

    def set(self, frac, prefix=None):
        CON.bar_set(frac, prefix)
        if not TRUECOLOR:
            p = prefix or self.prefix
            f = 0.0 if frac < 0 else (1.0 if frac > 1 else frac)
            filled = int(round(_BAR_W * f))
            sys.stdout.write(
                f"\r{p:<26} |{'█' * filled}{'░' * (_BAR_W - filled)}| "
                f"{int(f * 100):3d}%")
            sys.stdout.flush()

    def finish(self, ok=True, prefix=None):
        CON.bar_finish(ok, prefix)


CON = _Console(sys.stdout).start()
sys.stdout = CON


def getch_choice(valid):
    """Ждём одиночное нажатие одной из клавиш valid (без Enter). Иное — игнор."""
    try:
        import msvcrt
    except ImportError:
        # не Windows — читаем строкой
        while True:
            s = input().strip()
            if s in valid:
                return s
    while True:
        ch = msvcrt.getch()
        if ch in (b"\x00", b"\xe0"):   # спец-клавиши — съедаем второй байт
            msvcrt.getch()
            continue
        try:
            c = ch.decode("ascii")
        except UnicodeDecodeError:
            continue
        if c == "\x03":                # Ctrl+C
            raise KeyboardInterrupt
        if c in valid:
            return c
        # любая другая клавиша — ничего не делаем, ждём дальше


def read_line_visible(prompt):
    """Видимый построчный ввод. Подтверждение — Enter ИЛИ Tab."""
    try:
        import msvcrt
    except ImportError:
        return input(prompt)
    sys.stdout.write(prompt)
    sys.stdout.flush()
    buf = []
    while True:
        ch = msvcrt.getch()
        if ch in (b"\x00", b"\xe0"):       # спец-клавиши
            msvcrt.getch()
            continue
        if ch in (b"\r", b"\n", b"\t"):    # Enter или Tab — подтвердить
            sys.stdout.write("\n")
            sys.stdout.flush()
            return "".join(buf)
        if ch == b"\x08":                  # Backspace
            if buf:
                buf.pop()
                sys.stdout.write("\b \b")
                sys.stdout.flush()
            continue
        if ch == b"\x03":                  # Ctrl+C
            raise KeyboardInterrupt
        if ch == b"\x1b":                  # Esc — игнор внутри поля
            continue
        try:
            c = ch.decode("utf-8")
        except UnicodeDecodeError:
            try:
                c = (ch + msvcrt.getch()).decode("utf-8")
            except Exception:
                continue
        buf.append(c)
        sys.stdout.write(c)
        sys.stdout.flush()


def read_password(prompt="Пароль: "):
    """Скрытый ввод (эхо «*») через общий поток вывода. Подтверждение — Enter."""
    try:
        import msvcrt
    except ImportError:
        return getpass.getpass(prompt)
    sys.stdout.write(prompt)
    sys.stdout.flush()
    buf = []
    while True:
        ch = msvcrt.getch()
        if ch in (b"\x00", b"\xe0"):
            msvcrt.getch()
            continue
        if ch in (b"\r", b"\n"):
            sys.stdout.write("\n")
            sys.stdout.flush()
            return "".join(buf)
        if ch == b"\x08":
            if buf:
                buf.pop()
                sys.stdout.write("\b \b")
                sys.stdout.flush()
            continue
        if ch == b"\x03":
            raise KeyboardInterrupt
        try:
            c = ch.decode("utf-8")
        except UnicodeDecodeError:
            try:
                c = (ch + msvcrt.getch()).decode("utf-8")
            except Exception:
                continue
        buf.append(c)
        sys.stdout.write("*")
        sys.stdout.flush()


def wait_exit_or_again(ok=True):
    """Финальный выбор: ESC — выход, Enter — назад к выбору пользователя.
    Выше мигает SUCCESS (пока строка видна) — не задерживая ожидание."""
    CON.add_badge("✓  SUCCESS — отчёт сформирован" if ok else "✗  ERROR",
                  ok=ok)
    print(Fore.YELLOW
          + "  ESC — выход    |    Enter — назад к выбору пользователя"
          + Style.RESET_ALL)
    try:
        import msvcrt
    except ImportError:
        s = input("  Enter — ещё раз, q + Enter — выход: ").strip().lower()
        return "exit" if s == "q" else "again"
    while True:
        ch = msvcrt.getch()
        if ch in (b"\x00", b"\xe0"):
            msvcrt.getch()
            continue
        if ch == b"\x1b":
            return "exit"
        if ch in (b"\r", b"\n"):
            return "again"
        if ch == b"\x03":
            raise KeyboardInterrupt


def pause_exit(ok=False, label=None):
    """Показывает мигающий ERROR-бейдж и ждёт любую клавишу."""
    text = label or ("✗  ERROR — нажмите любую клавишу для выхода"
                     if not ok else "Нажмите любую клавишу для выхода")
    CON.add_badge(text, ok=ok)
    try:
        import msvcrt
    except ImportError:
        try:
            input()
        except Exception:
            pass
        return
    try:
        ch = msvcrt.getch()
        if ch in (b"\x00", b"\xe0"):
            msvcrt.getch()
    except Exception:
        pass


# ------------------------------------------------------------------ #
#  Доступность / подключение
# ------------------------------------------------------------------ #
def tcp_ok(host, port, timeout=2.0):
    try:
        with socket.create_connection((host, port), timeout=timeout):
            return True
    except OSError:
        return False


def resolve_endpoint(rec):
    """Возвращает (base_url, via, host, port) или None. Приоритет — локальная сеть."""
    h, p = rec.get("lan_ip", ""), rec.get("lan_port", 80)
    if h and tcp_ok(h, p):
        return (f"http://{h}:{p}", "локальная сеть", h, p)
    h, p = rec.get("rem_host", ""), rec.get("rem_port", 0)
    if h and p and tcp_ok(h, p):
        return (f"http://{h}:{p}", "удалённый порт", h, p)
    return None


def check_all(recs):
    print(Fore.YELLOW + "\nПроверка доступности регистраторов...\n" + Style.RESET_ALL)
    result = [None] * len(recs)
    n = max(len(recs), 1)
    bar = ProgressBar("Проверяю регистраторы").start()
    try:
        for i, rec in enumerate(recs):
            bar.set(i / n, f"Проверяю {rec['name']}")
            result[i] = resolve_endpoint(rec)
    finally:
        bar.finish(ok=any(v is not None for v in result), prefix="Доступность")
    return result


# ------------------------------------------------------------------ #
#  XML помощники (без привязки к namespace)
# ------------------------------------------------------------------ #
def _local(tag):
    return tag.split("}")[-1] if "}" in tag else tag


def findall_local(root, name):
    return [e for e in root.iter() if _local(e.tag) == name]


def first_text(elem, name, default=""):
    for e in elem.iter():
        if _local(e.tag) == name:
            return (e.text or "").strip()
    return default


# ------------------------------------------------------------------ #
#  Авторизация
# ------------------------------------------------------------------ #
def make_session(base, user, pw):
    """Пробуем Digest, затем Basic. Возвращаем (session, None) или (None, сообщение)."""
    last = ""
    for auth in (HTTPDigestAuth(user, pw), HTTPBasicAuth(user, pw)):
        s = requests.Session()
        s.auth = auth
        s.verify = False
        try:
            r = s.get(base + "/ISAPI/System/deviceInfo", timeout=HTTP_TIMEOUT)
        except requests.RequestException as e:
            last = f"сеть: {e}"
            continue
        if r.status_code == 200:
            return s, None
        if r.status_code == 401:
            last = "неверный логин или пароль"
        else:
            last = f"HTTP {r.status_code}"
    return None, last


# ------------------------------------------------------------------ #
#  Список пользователей
# ------------------------------------------------------------------ #
def get_users(session, base):
    r = session.get(base + "/ISAPI/Security/users", timeout=HTTP_TIMEOUT)
    r.raise_for_status()
    root = ET.fromstring(r.content)
    users = []
    for u in findall_local(root, "User"):
        name = first_text(u, "userName")
        level = first_text(u, "userLevel")
        if name:
            users.append({"name": name, "level": level})
    # уникальные, с сохранением порядка
    seen, uniq = set(), []
    for u in users:
        if u["name"] not in seen:
            seen.add(u["name"])
            uniq.append(u)
    return uniq


# ------------------------------------------------------------------ #
#  Карта "номер канала -> название камеры"
# ------------------------------------------------------------------ #
def get_channel_names(session, base):
    """Возвращает {'30': 'Проходная', ...}. Пробует несколько эндпоинтов ISAPI."""
    names = {}
    endpoints = [
        "/ISAPI/ContentMgmt/InputProxy/channels",
        "/ISAPI/ContentMgmt/InputProxy/channels/status",
        "/ISAPI/System/Video/inputs/channels",
    ]
    channel_tags = ("InputProxyChannel", "InputProxyChannelStatus",
                    "VideoInputChannel")
    for ep in endpoints:
        try:
            r = session.get(base + ep, timeout=HTTP_TIMEOUT)
        except requests.RequestException:
            continue
        if r.status_code != 200:
            continue
        try:
            root = ET.fromstring(r.content)
        except ET.ParseError:
            continue
        for tag in channel_tags:
            for ch in findall_local(root, tag):
                cid = first_text(ch, "id")
                nm = first_text(ch, "name")
                if cid and nm:
                    names.setdefault(str(cid).lstrip("Dd"), nm)
    return names


def cam_label(local_id, names):
    """'D30' + карта -> 'D30 — Проходная' (или 'D30', если имени нет)."""
    if not local_id:
        return ""
    digits = "".join(c for c in local_id if c.isdigit())
    nm = names.get(digits) or names.get(local_id)
    return f"{local_id} — {nm}" if nm else local_id


# ------------------------------------------------------------------ #
#  Поиск по журналу
# ------------------------------------------------------------------ #
def log_search(session, base, sub_type, start_iso, end_iso):
    """Тянем все записи заданного под-типа операции за период (с пагинацией)."""
    out = []
    pos = 0
    guard = 0
    while guard < 500:
        guard += 1
        sid = "{" + str(uuid.uuid4()).upper() + "}"
        body = (
            '<?xml version="1.0" encoding="utf-8"?>'
            "<CMSearchDescription>"
            f"<searchID>{sid}</searchID>"
            f"<metaId>log.std-cgi.com/Operation/{sub_type}</metaId>"
            "<timeSpanList><timeSpan>"
            f"<startTime>{start_iso}</startTime>"
            f"<endTime>{end_iso}</endTime>"
            "</timeSpan></timeSpanList>"
            "<maxResults>100</maxResults>"
            f"<searchResultPostion>{pos}</searchResultPostion>"
            "</CMSearchDescription>"
        )
        r = session.post(
            base + "/ISAPI/ContentMgmt/logSearch",
            data=body.encode("utf-8"),
            headers={"Content-Type": "application/xml"},
            timeout=HTTP_TIMEOUT,
        )
        if r.status_code != 200:
            break
        root = ET.fromstring(r.content)
        status = first_text(root, "responseStatusStrg", "OK")
        items = findall_local(root, "searchMatchItem")
        if not items:
            break
        for it in items:
            out.append({
                "t": first_text(it, "StartDateTime"),
                "user": first_text(it, "userName"),
                "ip": first_text(it, "ipAddress"),
                "id": first_text(it, "localId"),
                "info": first_text(it, "additionInformation"),
            })
        pos += len(items)
        if status != "MORE":
            break
    return out


def parse_dt(s):
    s = (s or "").strip().rstrip("Z")
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def fmt_duration(delta):
    if delta is None:
        return "—"
    total = int(delta.total_seconds())
    if total < 0:
        return "—"
    h, rem = divmod(total, 3600)
    m, sec = divmod(rem, 60)
    parts = []
    if h:
        parts.append(f"{h} ч")
    if m:
        parts.append(f"{m} мин")
    parts.append(f"{sec} сек")
    return " ".join(parts)


# Под-типы операций «действий» пользователя: код -> человекочитаемая метка
ACTION_SUBTYPES = {
    "remoteCtrlPtz": "PTZ (управление камерой)",
    "remotePlayByTime": "Воспроизведение архива",
    "remotePlayByFile": "Воспроизведение архива",
    "remotePlay": "Воспроизведение архива",
    "remoteGetByTime": "Скачивание архива",
    "remoteGetByFile": "Скачивание архива",
}


def collect_sessions(session, base, username, start_date, end_date, chan_names=None):
    """По дням тянем входы/выходы и действия (PTZ/воспроизведение) пользователя."""
    chan_names = chan_names or {}
    days = []
    d = start_date
    while d <= end_date:
        days.append(d)
        d += dt.timedelta(days=1)

    logins, logouts, actions = [], [], []
    total_steps = max(len(days), 1)
    ok_days = 0
    bar = ProgressBar("Анализ журнала").start()
    try:
        for i, day in enumerate(days):
            s_iso = day.strftime("%Y-%m-%dT00:00:00Z")
            e_iso = day.strftime("%Y-%m-%dT23:59:59Z")
            bar.set(i / total_steps, f"Журнал: {day.strftime('%d.%m.%Y')}")
            try:
                for ev in log_search(session, base, "remoteLogin", s_iso, e_iso):
                    if ev["user"] == username:
                        logins.append(ev)
                for ev in log_search(session, base, "remoteLogout", s_iso, e_iso):
                    if ev["user"] == username:
                        logouts.append(ev)
                for sub, label in ACTION_SUBTYPES.items():
                    for ev in log_search(session, base, sub, s_iso, e_iso):
                        if ev["user"] != username:
                            continue
                        t = parse_dt(ev["t"])
                        if not t:
                            continue
                        info = (ev.get("info") or "").replace("\n", " ").strip()
                        if len(info) > 160:
                            info = info[:160] + "…"
                        actions.append({
                            "t": t,
                            "type": label,
                            "cam": cam_label(ev["id"], chan_names),
                            "ip": ev.get("ip", ""),
                            "details": info,
                        })
                ok_days += 1
            except (requests.RequestException, ET.ParseError):
                pass
    finally:
        bar.finish(ok=(ok_days > 0), prefix="Анализ журнала")

    actions.sort(key=lambda a: a["t"])

    li = sorted([(parse_dt(e["t"]), e["ip"]) for e in logins if parse_dt(e["t"])],
                key=lambda x: x[0])
    lo = sorted([(parse_dt(e["t"]), e["ip"]) for e in logouts if parse_dt(e["t"])],
                key=lambda x: x[0])

    # Сопоставляем вход/выход СТРОГО по одному IP (один клиент).
    # Выход закрывает только текущую открытую сессию этого IP.
    # Если до следующего входа выхода нет — длительность неизвестна ("—"),
    # чтобы не склеивать разные сессии и не получать дни «активности».
    from collections import defaultdict
    ev = defaultdict(list)
    for t, ip in li:
        ev[ip].append((t, 0))      # 0 = login
    for t, ip in lo:
        ev[ip].append((t, 1))      # 1 = logout

    sessions = []
    for ip, lst in ev.items():
        # при равном времени логин обрабатываем раньше логаута (0 < 1)
        lst.sort(key=lambda x: (x[0], x[1]))
        open_login = None
        for t, typ in lst:
            if typ == 0:                       # вход
                if open_login is not None:
                    # предыдущий вход так и не закрылся выходом
                    sessions.append({"login": open_login, "logout": None,
                                     "dur": None, "ip": ip})
                open_login = t
            else:                              # выход
                if open_login is not None:
                    sessions.append({"login": open_login, "logout": t,
                                     "dur": t - open_login, "ip": ip})
                    open_login = None
                # выход без открытого входа — игнорируем (сирота)
        if open_login is not None:
            sessions.append({"login": open_login, "logout": None,
                             "dur": None, "ip": ip})

    sessions.sort(key=lambda s: s["login"])
    return sessions, actions


# ------------------------------------------------------------------ #
#  Excel
# ------------------------------------------------------------------ #
def _plural_cmd(n):
    n10, n100 = n % 10, n % 100
    if n10 == 1 and n100 != 11:
        w = "команда"
    elif 2 <= n10 <= 4 and not 12 <= n100 <= 14:
        w = "команды"
    else:
        w = "команд"
    return f"{n} {w}"


def _short_type(t):
    if t.startswith("PTZ"):
        return "PTZ"
    if "Воспроизвед" in t:
        return "Воспроизв."
    if "Скачив" in t:
        return "Скачив."
    return t


def _group_actions(acts, gap=180):
    """Схлопывает подряд идущие действия одного типа/камеры в один интервал."""
    acts = sorted(acts, key=lambda a: a["t"])
    groups = []
    for a in acts:
        if groups:
            g = groups[-1]
            if (g["type"] == a["type"] and g["cam"] == a["cam"]
                    and (a["t"] - g["end"]).total_seconds() <= gap):
                g["end"] = a["t"]
                g["count"] += 1
                continue
        groups.append({"type": a["type"], "cam": a["cam"],
                       "start": a["t"], "end": a["t"], "count": 1})
    return groups


def build_excel(path, meta, sessions, actions):
    wb = Workbook()
    ws = wb.active
    ws.title = "Входы и действия"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    sess_fill = PatternFill("solid", fgColor="DCE6F1")   # строка сессии
    act_fill = PatternFill("solid", fgColor="FFF7DC")    # строки действий
    head_font = Font(bold=True, color="FFFFFF", size=11)
    sess_font = Font(bold=True, color="1F3864")
    act_font = Font(italic=True, size=10, color="7A5C00")
    title_font = Font(bold=True, size=14, color="1F4E78")
    label_font = Font(bold=True, color="404040")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    ncols = 5
    ws.merge_cells("A1:E1")
    ws["A1"] = "Отчёт: входы и действия пользователя"
    ws["A1"].font = title_font
    ws["A1"].alignment = left

    n_ptz = sum(1 for a in actions if a["type"].startswith("PTZ"))
    n_play = sum(1 for a in actions if "Воспроизвед" in a["type"])
    n_dl = sum(1 for a in actions if "Скачив" in a["type"])

    info = [
        ("Пользователь:", meta["user"]),
        ("Регистратор:", meta["recorder"]),
        ("Подключение:", meta["via"]),
        ("Период:", meta["period"]),
        ("Сформировано:", meta["generated"]),
        ("Всего входов:", str(len(sessions))),
        ("С известной длительностью:", meta["known"]),
        ("Суммарное время (известное):", meta["total_dur"]),
        ("Действия (PTZ / воспр. / скач.):", f"{n_ptz} / {n_play} / {n_dl}"),
    ]
    row = 3
    for k, v in info:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        lc = ws.cell(row=row, column=1, value=k)
        lc.font = label_font
        lc.alignment = left
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        vc = ws.cell(row=row, column=3, value=v)
        vc.alignment = left
        row += 1

    header_row = row + 1
    headers = ["№", "Вход", "Выход", "Длительность", "IP-адрес"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = center
        cell.border = border

    # Привязка действий к сессиям по IP (сессия с последним входом <= времени)
    from collections import defaultdict
    by_ip = defaultdict(list)
    for idx, s in enumerate(sessions):
        s["acts"] = []
        by_ip[s["ip"]].append(idx)
    unattached = []
    for a in actions:
        chosen = None
        for idx in by_ip.get(a["ip"], []):      # индексы в порядке возрастания входа
            if sessions[idx]["login"] <= a["t"]:
                chosen = idx
            else:
                break
        if chosen is None:
            unattached.append(a)
        else:
            sessions[chosen]["acts"].append(a)

    def write_action_group(r, g):
        cam_txt = g["cam"] or "—"
        if g["count"] > 1:
            cam_txt += f"   ({_plural_cmd(g['count'])})"
        if g["start"] == g["end"]:
            tr = g["start"].strftime("%H:%M:%S")
        else:
            tr = (g["start"].strftime("%H:%M:%S") + "–"
                  + g["end"].strftime("%H:%M:%S"))
        ws.cell(row=r, column=1, value="")
        ws.cell(row=r, column=2, value="     ↳ " + tr)
        ws.cell(row=r, column=3, value=_short_type(g["type"]))
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        ws.cell(row=r, column=4, value=cam_txt)
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = act_fill
            cell.font = act_font
            cell.border = border
            cell.alignment = left if c in (2, 4) else center

    r = header_row + 1
    for i, s in enumerate(sessions, 1):
        has_out = s["logout"] is not None
        vals = [
            i,
            s["login"].strftime("%d.%m.%Y %H:%M:%S"),
            s["logout"].strftime("%d.%m.%Y %H:%M:%S") if has_out else "—",
            fmt_duration(s["dur"]),
            s["ip"] or "—",
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = center
            cell.fill = sess_fill
            cell.font = sess_font if c == 1 else Font(color="1F3864")
            if not has_out and c in (3, 4):
                cell.font = Font(color="A6A6A6")
        r += 1
        for g in _group_actions(s["acts"]):
            write_action_group(r, g)
            r += 1

    if not sessions:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        c = ws.cell(row=r, column=1,
                    value="За выбранный период входов пользователя не найдено")
        c.alignment = center
        c.font = Font(italic=True, color="C00000")
        r += 1

    if unattached:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        c = ws.cell(row=r, column=1,
                    value="Действия вне распознанных сессий (вход не найден):")
        c.font = Font(bold=True, italic=True, color="7A5C00")
        c.alignment = left
        r += 1
        for g in _group_actions(unattached):
            # для внесессионных показываем дату+время
            cam_txt = g["cam"] or "—"
            if g["count"] > 1:
                cam_txt += f"   ({_plural_cmd(g['count'])})"
            tr = g["start"].strftime("%d.%m.%Y %H:%M:%S")
            if g["start"] != g["end"]:
                tr += "–" + g["end"].strftime("%H:%M:%S")
            ws.cell(row=r, column=2, value="     ↳ " + tr)
            ws.cell(row=r, column=3, value=_short_type(g["type"]))
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
            ws.cell(row=r, column=4, value=cam_txt)
            for c in range(1, ncols + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = act_fill
                cell.font = act_font
                cell.border = border
                cell.alignment = left if c in (2, 4) else center
            r += 1

    for i, w in enumerate([6, 24, 22, 20, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    wb.save(path)


# ------------------------------------------------------------------ #
#  Ввод дат
# ------------------------------------------------------------------ #
def ask_date(prompt, default=None):
    while True:
        raw = read_line_visible(prompt).strip()
        if not raw and default is not None:
            return default
        try:
            return dt.datetime.strptime(raw, "%d.%m.%Y").date()
        except ValueError:
            print(Fore.RED + "  Неверный формат. Пример: 17.10.2026" + Style.RESET_ALL)


# ------------------------------------------------------------------ #
#  Настройки регистраторов (меню по «0»)
# ------------------------------------------------------------------ #
def press_any():
    try:
        import msvcrt
        print(Fore.YELLOW + "  — нажмите любую клавишу —" + Style.RESET_ALL)
        ch = msvcrt.getch()
        if ch in (b"\x00", b"\xe0"):
            msvcrt.getch()
    except Exception:
        try:
            input()
        except Exception:
            pass


def _edit_recorder_dialog(rec):
    """Диалог создания/редактирования. Пустой ввод = оставить текущее."""
    cur = rec or {"name": "", "lan_ip": "", "lan_port": 80,
                  "rem_host": "", "rem_port": 0}

    def ask(label, val):
        hint = f" [{val}]" if val not in (None, "", 0) else ""
        return read_line_visible(f"  {label}{hint}: ").strip()

    print()
    print(Fore.CYAN + ("  Новый регистратор" if rec is None
                       else f"  Редактирование: {cur['name']}") + Style.RESET_ALL)
    name = ask("Название", cur["name"]) or cur["name"]
    lan_ip = ask("Внутренний IP (локальная сеть)", cur["lan_ip"]) or cur["lan_ip"]
    lp = ask("Внутренний порт (обычно 80)", cur["lan_port"])
    lan_port = int(lp) if lp.isdigit() else cur["lan_port"]
    rem_host = ask("Внешний IP/хост (для удалёнки)", cur["rem_host"])
    rem_host = rem_host if rem_host else cur["rem_host"]
    rp = ask("Внешний порт", cur["rem_port"])
    rem_port = int(rp) if rp.isdigit() else cur["rem_port"]

    if not lan_ip and not rem_host:
        print(Fore.RED + "  Нужен хотя бы внутренний IP или внешний хост."
              + Style.RESET_ALL)
        press_any()
        return None
    if not name:
        name = lan_ip or rem_host or "регистратор"
    return {"name": name, "lan_ip": lan_ip, "lan_port": lan_port,
            "rem_host": rem_host, "rem_port": rem_port}


def settings_menu(recs):
    """Меню настроек: изменить / добавить / удалить регистраторы, сохранить в exe."""
    while True:
        cls()
        banner("Настройки регистраторов")
        if not recs:
            print(Fore.YELLOW + "  Список пуст — добавьте регистратор (A)."
                  + Style.RESET_ALL)
        for i, r in enumerate(recs, 1):
            rem = (f"{r['rem_host']}:{r['rem_port']}"
                   if r["rem_host"] and r["rem_port"] else "—")
            print(f"  {Style.BRIGHT}{i}{Style.NORMAL}. {r['name']}")
            print(f"       внутренний: {r['lan_ip']}:{r['lan_port']}")
            print(f"       внешний:    {rem}")
        print()
        print("  [цифра] изменить   [A] добавить   [D] удалить")
        print("  [S] сохранить и назад   [Q] назад без сохранения")
        ans = read_line_visible(Fore.YELLOW + "  Выбор: "
                                + Style.RESET_ALL).strip().lower()

        if ans in ("q", ""):
            return
        if ans == "s":
            ok, err = save_recorders(recs)
            if ok:
                CON.add_badge("✓  SUCCESS — настройки сохранены в exe", ok=True)
            else:
                CON.add_badge(f"✗  ERROR — не сохранено: {err}", ok=False)
            press_any()
            return
        if ans == "a":
            r = _edit_recorder_dialog(None)
            if r:
                recs.append(r)
            continue
        if ans == "d":
            num = read_line_visible("  Номер для удаления: ").strip()
            if num.isdigit() and 1 <= int(num) <= len(recs):
                removed = recs.pop(int(num) - 1)
                print(Fore.GREEN + f"  Удалён: {removed['name']}"
                      + Style.RESET_ALL)
                press_any()
            continue
        if ans.isdigit() and 1 <= int(ans) <= len(recs):
            r = _edit_recorder_dialog(recs[int(ans) - 1])
            if r:
                recs[int(ans) - 1] = r
            continue


# ------------------------------------------------------------------ #
#  Главный сценарий
# ------------------------------------------------------------------ #
def main():
    cls()
    try:
        import ctypes
        ctypes.windll.kernel32.SetConsoleTitleW("HIK-Report")
    except Exception:
        pass
    banner("HIK-Report — отчёт по входам (Hikvision NVR)")

    # Шаг 1-2 — выбор регистратора (0 — настройки)
    recs = load_recorders()
    while True:
        avail = check_all(recs) if recs else []
        print()
        banner("Выберите регистратор")
        if not recs:
            print(Fore.YELLOW + "  Список пуст. Нажмите 0, чтобы добавить "
                  "регистратор." + Style.RESET_ALL)
        for i, rec in enumerate(recs, 1):
            ep = avail[i - 1]
            if ep:
                tag = Fore.GREEN + f"[доступен — {ep[1]}]" + Style.RESET_ALL
            else:
                tag = Fore.RED + Style.BRIGHT + "[НЕДОСТУПЕН]" + Style.RESET_ALL
            print(f"  {Style.BRIGHT}{i}{Style.NORMAL}. {rec['name']}   {tag}")
        print(f"  {Style.BRIGHT}0{Style.NORMAL}. Настройки "
              "(изменить / добавить регистраторы)")
        print(Fore.YELLOW + "\n  Выберите цифру регистратора или 0 — настройки"
              + Style.RESET_ALL)

        valid = {"0"} | {str(i) for i in range(1, min(len(recs), 9) + 1)}
        choice = getch_choice(valid)
        if choice == "0":
            settings_menu(recs)
            recs = load_recorders()
            continue
        idx = int(choice) - 1
        rec = recs[idx]
        ep = avail[idx]
        if not ep:
            print(Fore.RED + f"\n  Регистратор {rec['name']} недоступен ни "
                  f"локально, ни по удалённому порту." + Style.RESET_ALL)
            CON.add_badge("✗  ERROR — регистратор недоступен, выберите другой",
                          ok=False)
            press_any()
            continue
        base_url, via = ep[0], ep[1]
        break

    # Шаг 4 — подключение c прогрессбаром
    print()
    steps = ["Определение адреса", "TCP-соединение",
             "Проверка ISAPI", "Канал готов"]
    bar = ProgressBar("Подключение").start()
    try:
        for i, st in enumerate(steps):
            bar.set(i / len(steps), st)
            time.sleep(0.45)
    finally:
        bar.finish(ok=True, prefix="Подключение")
    print(Fore.GREEN + f"Подключено к {rec['name']} ({via}): {base_url}"
          + Style.RESET_ALL)

    # логин/пароль в терминале
    session = None
    for attempt in range(3):
        print()
        user = read_line_visible("Логин: ").strip()
        pw = read_password("Пароль: ")
        print(Fore.YELLOW + "Проверка учётных данных…" + Style.RESET_ALL)
        session, err = make_session(base_url, user, pw)
        if session:
            print(Fore.GREEN + "Авторизация успешна." + Style.RESET_ALL)
            break
        print(Fore.RED + f"Ошибка входа: {err}" + Style.RESET_ALL)
    if not session:
        print(Fore.RED + "Не удалось авторизоваться. Выход." + Style.RESET_ALL)
        pause_exit()
        return

    # Шаг 5 — пользователи
    print(Fore.YELLOW + "\nСчитываю список пользователей…" + Style.RESET_ALL)
    try:
        users = get_users(session, base_url)
    except Exception as e:
        print(Fore.RED + f"Не удалось получить список пользователей: {e}"
              + Style.RESET_ALL)
        pause_exit()
        return
    if not users:
        print(Fore.RED + "Пользователи не найдены." + Style.RESET_ALL)
        pause_exit()
        return

    # Карта "канал -> название камеры" (один раз на подключение)
    print(Fore.YELLOW + "Считываю названия камер…" + Style.RESET_ALL)
    try:
        chan_names = get_channel_names(session, base_url)
    except Exception:
        chan_names = {}
    print(Fore.GREEN + f"  Камер с названиями: {len(chan_names)}" + Style.RESET_ALL)

    # Цикл: выбор пользователя -> даты -> отчёт. ESC — выход, Enter — заново.
    while True:
        print()
        banner("Пользователи на регистраторе")
        for i, u in enumerate(users, 1):
            lvl = f"  ({u['level']})" if u["level"] else ""
            print(f"  {Style.BRIGHT}{i:>2}{Style.NORMAL}. {u['name']}{lvl}")

        if len(users) <= 9:
            print(Fore.YELLOW + "\n  Нажмите цифру пользователя"
                  + Style.RESET_ALL)
            valid = {str(i) for i in range(1, len(users) + 1)}
            sel = int(getch_choice(valid))
        else:
            while True:
                raw = read_line_visible("\n  Введите номер пользователя "
                                        "и Enter: ").strip()
                if raw.isdigit() and 1 <= int(raw) <= len(users):
                    sel = int(raw)
                    break
                print(Fore.RED + "  Неверный номер." + Style.RESET_ALL)
        username = users[sel - 1]["name"]
        print(Fore.GREEN + f"  Выбран пользователь: {username}" + Style.RESET_ALL)

        # Шаг 6 — даты
        print()
        banner("Диапазон дат")
        today = dt.date.today()
        start_date = ask_date("Выберите дату начала в формате dd.mm.yyyy: ")
        end_date = ask_date(
            f"Выберите дату окончания dd.mm.yyyy (Enter = сегодня {today.strftime('%d.%m.%Y')}): ",
            default=today,
        )
        if end_date < start_date:
            start_date, end_date = end_date, start_date

        # Шаг 7 — сбор данных
        print()
        banner("Анализ журнала")
        sessions, actions = collect_sessions(session, base_url, username,
                                             start_date, end_date, chan_names)

        total = dt.timedelta()
        known = 0
        for s in sessions:
            if s["dur"]:
                total += s["dur"]
                known += 1

        print(Fore.GREEN + f"\nНайдено входов: {len(sessions)}"
              f"  (с известной длительностью: {known}); "
              f"действий: {len(actions)}" + Style.RESET_ALL)

        # Шаг 8 — Excel
        safe_user = "".join(ch for ch in username if ch.isalnum() or ch in "-_")
        fname = (f"otchet_{safe_user}_{rec['name'].replace('.', '-')}_"
                 f"{start_date.strftime('%d%m%Y')}-{end_date.strftime('%d%m%Y')}.xlsx")
        out_path = os.path.join(output_dir(), fname)

        meta = {
            "user": username,
            "recorder": rec["name"],
            "via": via,
            "period": f"{start_date.strftime('%d.%m.%Y')} — {end_date.strftime('%d.%m.%Y')}",
            "generated": dt.datetime.now().strftime("%d.%m.%Y %H:%M"),
            "total_dur": fmt_duration(total) if total.total_seconds() else "—",
            "known": f"{known} из {len(sessions)}",
        }
        print(Fore.YELLOW + "Формирую Excel…" + Style.RESET_ALL)
        build_excel(out_path, meta, sessions, actions)
        print(Fore.GREEN + Style.BRIGHT + f"\nГотово! Файл: {out_path}"
              + Style.RESET_ALL)

        try:
            os.startfile(out_path)   # откроет Excel автоматически (Windows)
        except Exception:
            pass

        if wait_exit_or_again() == "exit":
            break


def selftest():
    """Проверка, что все зависимости на месте (для сборки .exe). Без диалога."""
    import requests as _rq
    import openpyxl as _ox
    import colorama as _cl
    # собираем пустой Excel в память — проверяем движок openpyxl
    build_excel(os.path.join(output_dir(), "_selftest_tmp.xlsx"),
                {"user": "test", "recorder": "-", "via": "-", "period": "-",
                 "generated": "-", "total_dur": "—", "known": "0 из 0"}, [], [])
    try:
        os.remove(os.path.join(output_dir(), "_selftest_tmp.xlsx"))
    except Exception:
        pass
    frozen = "да" if getattr(sys, "frozen", False) else "нет"
    print(f"SELFTEST OK | requests {_rq.__version__} | openpyxl {_ox.__version__} "
          f"| colorama {_cl.__version__} | frozen={frozen}")


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
        sys.exit(0)
    if "--adstest" in sys.argv:
        ok, err = save_recorders(load_recorders())
        rd = load_recorders()
        print("ADSTEST write:", ok, err, "| read_back:", len(rd))
        sys.exit(0)
    try:
        main()
    except KeyboardInterrupt:
        print(Fore.RED + "\nПрервано пользователем." + Style.RESET_ALL)
    except Exception as e:
        print(Fore.RED + f"\nНеожиданная ошибка: {e}" + Style.RESET_ALL)
        pause_exit()
    finally:
        try:
            CON.stop()
        except Exception:
            pass
